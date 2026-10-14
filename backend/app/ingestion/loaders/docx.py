"""
DOCX document loader — extracts text from Word documents.
"""

from docx import Document as DocxDocument
from app.ingestion.loaders.base import BaseLoader, Document
from app.models.exceptions import DocumentLoadError


class DOCXLoader(BaseLoader):
    """Load DOCX (.docx) files and extract text content."""

    SUPPORTED_EXTENSIONS = [".docx"]

    def load(self, filepath: str) -> list[Document]:
        try:
            doc = DocxDocument(filepath)
            paragraphs = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)

            content = "\n\n".join(paragraphs)

            if not content.strip():
                return []

            metadata = self._base_metadata(filepath)

            # Extract core properties
            if doc.core_properties.title:
                metadata["title"] = doc.core_properties.title
            if doc.core_properties.author:
                metadata["author"] = doc.core_properties.author

            return [Document(content=content, metadata=metadata)]
        except Exception as e:
            raise DocumentLoadError(filepath, str(e))
